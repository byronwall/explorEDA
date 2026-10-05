from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
root=Path(__file__).resolve().parent.parent
doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.62);sec.left_margin=Inches(.72);sec.right_margin=Inches(.72)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.25)
styles=doc.styles
for name in ['Normal','Body Text']:
    st=styles[name];st.font.name='Calibri';st.font.size=Pt(10.5)
    st.font.color.rgb=RGBColor.from_string('263547');st.paragraph_format.line_spacing=1.08;st.paragraph_format.space_after=Pt(6)
for name,size in [('Title',34),('Subtitle',17),('Heading 1',20),('Heading 2',15),('Heading 3',12)]:
    st=styles[name];st.font.name='Calibri';st.font.size=Pt(size);st.font.color.rgb=RGBColor.from_string('153B63')
    st.font.bold=name!='Subtitle';st.paragraph_format.space_before=Pt(8);st.paragraph_format.space_after=Pt(8)
    st.paragraph_format.keep_with_next=True
code=styles.add_style('Outline Code',1);code.font.name='DejaVu Sans Mono';code.font.size=Pt(8.5)
code.font.color.rgb=RGBColor.from_string('16324F');code.paragraph_format.line_spacing=1.08
code.paragraph_format.space_after=Pt(0);code.paragraph_format.left_indent=Inches(.10);code.paragraph_format.right_indent=Inches(.06)
small=styles.add_style('Review Small',1);small.font.name='Calibri';small.font.size=Pt(8.5);small.font.color.rgb=RGBColor.from_string('607286')
header=sec.header.paragraphs[0];header.style=small;header.text='EXPLOR EDA  /  OUTLINE                                        FINAL LANGUAGE REVIEW 03'
footer=sec.footer.paragraphs[0];footer.style=small;footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('3 OCTOBER 2026   ·   ')
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
# Document properties deliberately contain no user paths or private runtime metadata.
doc.core_properties.title='Outline — Final language review';doc.core_properties.subject='Flat explorEDA chart DSL with calculations and chart-owned filters'
doc.core_properties.author='';doc.core_properties.keywords='explorEDA, DSL, calculations, filters, design review'

def inline(p,text,size=None,bold=False):
    for token in re.split(r'(\*\*.*?\*\*|`[^`]*`)',text):
        if not token:
            continue
        if token.startswith('**'):
            inline(p,token[2:-2],size,bold=True)
            continue
        is_code=token.startswith('`')
        r=p.add_run(token[1:-1] if is_code else token)
        r.bold=bold
        if is_code:
            r.font.name='DejaVu Sans Mono'
            r.font.size=Pt(size or 8.7)
            r.font.color.rgb=RGBColor.from_string('174C74')
        elif size:
            r.font.size=Pt(size)

def shade(p,color='F1F5F9'):
    pr=p._p.get_or_add_pPr();s=OxmlElement('w:shd');s.set(qn('w:fill'),color);pr.append(s)

def table(rows):
    t=doc.add_table(rows=1, cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    widths=[2.38,4.68] if len(rows[0])==2 else [2.1]*len(rows[0])
    if rows[0][0]=='Ref':widths=[.58,6.48]
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for ri,row in enumerate(rows):
        cells=t.rows[0].cells if ri==0 else t.add_row().cells
        for ci,text in enumerate(row):
            c=cells[ci];c.width=Inches(widths[ci]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=c._tc.get_or_add_tcPr();m=OxmlElement('w:tcMar')
            for edge in ['top','bottom','left','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'75' if edge in ['top','bottom'] else '90');e.set(qn('w:type'),'dxa');m.append(e)
            tcpr.append(m);sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5EEF7' if ri==0 else ('F5F7FA' if ri%2==0 else 'FFFFFF'));tcpr.append(sh)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.04
            inline(p,text,9.2)
            if ri==0:
                for r in p.runs:r.bold=True
        pr=t.rows[ri]._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');pr.append(no)
        if ri==0:repeat=OxmlElement('w:tblHeader');pr.append(repeat)
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

lines=(root/'REVIEW.md').read_text().splitlines();i=0;code_open=False;code_lines=[]
while i<len(lines):
    line=lines[i]
    if line.startswith('```'):
        if not code_open:code_open=True;code_lines=[]
        else:
            for index,l in enumerate(code_lines):
                p=doc.add_paragraph(style='Outline Code');p.add_run(l or ' ');shade(p)
                p.paragraph_format.keep_with_next=index<len(code_lines)-1
                p.paragraph_format.keep_together=True
            spacer=doc.add_paragraph();spacer.paragraph_format.space_after=Pt(2);spacer.paragraph_format.space_before=Pt(0);spacer.paragraph_format.line_spacing=Pt(2)
            code_open=False
        i+=1;continue
    if code_open:code_lines.append(line);i+=1;continue
    if line=='<!-- PAGE -->':doc.add_page_break();i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].startswith('|'):
            # escaped Markdown pipes don't occur in this review's tables.
            parts=[x.strip() for x in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r'[:\- ]+',x) for x in parts):rows.append(parts)
            i+=1
        table(rows);continue
    if not line.strip():i+=1;continue
    if line.startswith('# '):doc.add_paragraph(line[2:],style='Title')
    elif line.startswith('## '):doc.add_paragraph(line[3:],style='Subtitle' if line=='## Final language review' else 'Heading 1')
    elif line.startswith('### '):doc.add_paragraph(line[4:],style='Heading 2')
    else:inline(doc.add_paragraph(),line)
    i+=1
# Hyperlink reference IDs in the source table to the primary source entries.
register=(root/'SOURCES.md').read_text()
links={}
for ref,url in re.findall(r'\*\*(R\d+):\*\*.*?\]\((https?://[^)]+)\)',register):links[ref]=url
for t in doc.tables:
    if t.cell(0,0).text!='Ref':continue
    for row in t.rows[1:]:
        ref=row.cells[0].text
        if ref not in links:continue
        p=row.cells[0].paragraphs[0];p.clear()
        h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(links[ref],RT.HYPERLINK,is_external=True))
        r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'1B5894');pr.append(c);r.append(pr);txt=OxmlElement('w:t');txt.text=ref;r.append(txt);h.append(r);p._p.append(h)
path=root/'outline-final-review.docx';doc.save(path);print(path)
